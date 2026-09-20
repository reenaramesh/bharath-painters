import json
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

from quotations.models import ApartmentCommunity


RESIDENTIAL_CATEGORIES = {
    "apartment building",
    "apartment complex",
    "apartment rental agency",
    "condominium complex",
    "furnished apartment building",
    "gated community",
    "homeowners' association",
    "housing complex",
    "housing cooperative",
    "housing development",
    "housing society",
    "military residence",
    "multi-unit residential building",
    "residents association",
    "retirement community",
    "serviced accommodation",
    "serviced apartment",
    "short term apartment rental agency",
    "student dormitory",
    "villa",
}


class Command(BaseCommand):
    help = "Import Bengaluru apartment communities from an Excel or Google Places JSON file."

    def add_arguments(self, parser):
        parser.add_argument("source", help="Path to the .xlsx or .json apartment directory")

    def handle(self, *args, **options):
        source = Path(options["source"])
        if source.suffix.lower() == ".json":
            records = self.json_records(source)
        elif source.suffix.lower() == ".xlsx":
            records = self.workbook_records(source)
        else:
            raise CommandError("Only .xlsx and .json apartment directories are supported.")

        created = updated = skipped = excluded = 0
        seen = set()
        for record in records:
            if record is None:
                excluded += 1
                continue
            name = record["name"]
            pincode = record["pincode"]
            key = (name.casefold(), pincode)
            if key in seen:
                skipped += 1
                continue
            seen.add(key)

            community = ApartmentCommunity.objects.filter(
                name__iexact=name, pincode=pincode,
            ).first()
            values = {
                "zone": record["zone"],
                "locality": record["locality"],
                "map_url": record["map_url"],
                "latitude": record["latitude"],
                "longitude": record["longitude"],
                "is_active": True,
            }
            if community:
                for field, value in values.items():
                    if value not in (None, "") or field == "is_active":
                        setattr(community, field, value)
                community.save()
                updated += 1
            else:
                ApartmentCommunity.objects.create(name=name, pincode=pincode, **values)
                created += 1

        details = (
            f"Apartment directory imported: {created} created, {updated} updated, "
            f"{skipped} duplicate rows skipped"
        )
        if excluded:
            details += f", {excluded} unrelated or incomplete places excluded"
        self.stdout.write(self.style.SUCCESS(details + "."))

    def workbook_records(self, source):
        try:
            sheet = load_workbook(source, read_only=True, data_only=True).active
        except (OSError, ValueError) as exc:
            raise CommandError(f"Workbook could not be opened: {exc}") from exc

        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows, ())]
        indexes = {name: position for position, name in enumerate(headers)}
        required = {"Apartment Name", "Zone", "Locality", "Pin Code"}
        if not required.issubset(indexes):
            raise CommandError("Workbook does not contain the expected apartment columns.")

        for row in rows:
            name = self.cell(row, indexes, "Apartment Name")
            if not name:
                continue
            latitude, longitude = self.coordinates(
                self.cell(row, indexes, "Latitude, Longitude"),
            )
            yield {
                "name": name,
                "zone": self.cell(row, indexes, "Zone"),
                "locality": self.cell(row, indexes, "Locality"),
                "pincode": self.pincode(self.cell(row, indexes, "Pin Code")),
                "map_url": self.cell(row, indexes, "Location Map (web link)"),
                "latitude": latitude,
                "longitude": longitude,
            }

    def json_records(self, source):
        try:
            payload = json.loads(source.read_text(encoding="utf-8-sig"))
        except (OSError, ValueError) as exc:
            raise CommandError(f"JSON file could not be opened: {exc}") from exc
        rows = payload if isinstance(payload, list) else payload.get("items") or payload.get("results")
        if not isinstance(rows, list):
            raise CommandError("JSON must be a list or contain an items/results list.")

        for row in rows:
            if not isinstance(row, dict):
                yield None
                continue
            name = str(row.get("title") or row.get("name") or "").strip()
            category = str(row.get("categoryName") or row.get("category") or "").strip().casefold()
            if not name or category not in RESIDENTIAL_CATEGORIES:
                yield None
                continue
            address = str(row.get("address") or "").strip()
            location = row.get("location") if isinstance(row.get("location"), dict) else {}
            yield {
                "name": name,
                "zone": str(row.get("state") or "").strip(),
                "locality": str(row.get("neighborhood") or row.get("city") or "").strip(),
                "pincode": self.pincode(row.get("postalCode") or self.address_pincode(address)),
                "map_url": str(row.get("url") or "").strip(),
                "latitude": self.decimal(location.get("lat") or row.get("latitude")),
                "longitude": self.decimal(location.get("lng") or row.get("longitude")),
            }

    @staticmethod
    def cell(row, indexes, name):
        position = indexes.get(name)
        if position is None or position >= len(row):
            return ""
        value = row[position]
        return str(value).strip() if value is not None else ""

    @staticmethod
    def pincode(value):
        value = str(value or "").strip()
        return value[:-2] if value.endswith(".0") else value

    @staticmethod
    def address_pincode(address):
        match = re.search(r"\b[1-9][0-9]{5}\b", address)
        return match.group(0) if match else ""

    @staticmethod
    def decimal(value):
        if value in (None, ""):
            return None
        try:
            return Decimal(str(value))
        except (InvalidOperation, ValueError):
            return None

    @classmethod
    def coordinates(cls, value):
        parts = [part.strip() for part in str(value or "").split(",")]
        if len(parts) != 2:
            return None, None
        return cls.decimal(parts[0]), cls.decimal(parts[1])
