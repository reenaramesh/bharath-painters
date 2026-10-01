"""CSV/XLSX import for editable quotation and apartment master data."""

import csv
import io
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.db.models import Q
from openpyxl import load_workbook

from accounts.models import BharathUser
from .models import (
    ApartmentCommunity, Area, MeasurementSurfaceType, ServiceCategory,
    PaintType, WorkDescription, PaintBrand, Unit,
)


MASTER_MODELS = {
    "areas": Area,
    "measurement-surface-types": MeasurementSurfaceType,
    "service-categories": ServiceCategory,
    "paint-types": PaintType,
    "work-descriptions": WorkDescription,
    "brands": PaintBrand,
    "units": Unit,
    "apartments": ApartmentCommunity,
}
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000


def header_key(value):
    return str(value or "").strip().lower().replace(" ", "_").replace("-", "_")


def read_master_rows(upload):
    if not upload:
        raise ValueError("Choose a CSV or XLSX file.")
    if upload.size > MAX_UPLOAD_BYTES:
        raise ValueError("The upload must be 5 MB or smaller.")
    filename = upload.name.lower()
    if filename.endswith(".csv"):
        try:
            reader = csv.DictReader(io.StringIO(upload.read().decode("utf-8-sig")))
            rows = [{header_key(key): str(value or "").strip() for key, value in row.items() if key}
                    for row in reader]
        except UnicodeDecodeError as error:
            raise ValueError("Save the CSV file in UTF-8 format.") from error
    elif filename.endswith(".xlsx"):
        try:
            workbook = load_workbook(upload, read_only=True, data_only=True)
            values = workbook.active.iter_rows(values_only=True)
            headers = [header_key(value) for value in next(values, ())]
            rows = [dict(zip(headers, (str(value).strip() if value is not None else "" for value in row)))
                    for row in values]
            workbook.close()
        except (OSError, ValueError, KeyError) as error:
            raise ValueError("The Excel file could not be read.") from error
    else:
        raise ValueError("Only CSV and XLSX files are supported.")
    if len(rows) > MAX_ROWS:
        raise ValueError("Upload at most 5,000 rows at a time.")
    if not rows:
        raise ValueError("The file has no data rows.")
    return rows


def row_value(row, *names):
    return next((row[name] for name in names if row.get(name)), "")


def import_master_rows(section, upload, user):
    model = MASTER_MODELS[section]
    rows = read_master_rows(upload)
    created = updated = skipped = 0
    errors = []
    for number, row in enumerate(rows, 2):
        try:
            with transaction.atomic():
                result = import_master_row(model, row, user)
            if result == "created":
                created += 1
            elif result == "updated":
                updated += 1
            else:
                skipped += 1
        except (ValueError, InvalidOperation) as error:
            if len(errors) < 30:
                errors.append({"row": number, "message": str(error)})
    return {"created": created, "updated": updated, "skipped": skipped, "errors": errors,
            "error_count": len(rows) - created - updated - skipped}


def next_sort_order(model):
    return (model.objects.aggregate(top=models.Max("sort_order"))["top"] or 0) + 1


def import_master_row(model, row, user):
    name = row_value(row, "name", "apartment_name")
    limit = 200 if model is ApartmentCommunity else 150
    if not name or len(name) > limit:
        raise ValueError(f"Name is required and must be at most {limit} characters.")

    if model is ApartmentCommunity:
        pincode = row_value(row, "pincode", "pin_code", "postal_code")
        zone = row_value(row, "zone", "state")
        locality = row_value(row, "locality", "neighborhood")
        if len(pincode) > 10 or len(zone) > 120 or len(locality) > 160:
            raise ValueError("PIN code, zone or locality is too long.")
        item = model.objects.filter(name__iexact=name, pincode=pincode).first()
        if item:
            item.zone, item.locality, item.is_active = zone, locality, True
            item.save(update_fields=["zone", "locality", "is_active", "updated_at"])
            return "updated"
        model.objects.create(name=name, pincode=pincode, zone=zone, locality=locality, sort_order=next_sort_order(model))
        return "created"

    category = None
    if model in (PaintType, WorkDescription):
        category_name = row_value(row, "service_category", "type_of_service", "category")
        if not category_name:
            raise ValueError("Type of Service is required for this section.")
        visible = ServiceCategory.objects.filter(is_active=True).filter(
            Q(created_by__isnull=True) | Q(created_by__role=BharathUser.Roles.ADMIN) | Q(created_by=user)
        )
        category = visible.filter(name__iexact=category_name).first()
        if not category:
            raise ValueError(f"Type of Service '{category_name}' does not exist in Master Data.")

    filters = {"name__iexact": name}
    values = {"name": name, "is_active": True}
    if category:
        filters["service_category"] = category
        values["service_category"] = category
    if model is PaintType:
        values["key_features"] = row_value(row, "key_features", "features")
        price = row_value(row, "default_price", "price")
        if price:
            amount = Decimal(price.replace(",", ""))
            if amount < 0 or amount >= Decimal("10000000000"):
                raise ValueError("Default price is outside the allowed range.")
            values["default_price"] = amount
        else:
            values["default_price"] = None

    owned = model.objects.filter(Q(created_by__isnull=True) | Q(created_by__role=BharathUser.Roles.ADMIN)) if user.role == BharathUser.Roles.ADMIN else model.objects.filter(created_by=user)
    item = owned.filter(**filters).first()
    if item:
        for field, value in values.items():
            setattr(item, field, value)
        item.save(update_fields=[*values, "updated_at"])
        return "updated"
    if user.role == BharathUser.Roles.CONTRACTOR and model.objects.filter(
        Q(created_by__isnull=True) | Q(created_by__role=BharathUser.Roles.ADMIN), **filters
    ).exists():
        return "skipped"
    model.objects.create(created_by=None if user.role == BharathUser.Roles.ADMIN else user, sort_order=next_sort_order(model), **values)
    return "created"
